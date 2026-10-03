import io

import pytest
from fastapi.testclient import TestClient

from tracetool.app import create_app
from tracetool.db import Database


@pytest.fixture
def db():
    d = Database(":memory:")
    yield d
    d.close()


@pytest.fixture
def client(db):
    return TestClient(create_app(db))


def make_xlsx(sheets: dict[str, list[list]], merges: dict[str, list[str]] | None = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for rng in (merges or {}).get(name, []):
            ws.merge_cells(rng)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_csv(rows: list[list[str]], encoding: str = "utf-8", delimiter: str = ",") -> bytes:
    import csv

    buf = io.StringIO()
    csv.writer(buf, delimiter=delimiter, lineterminator="\r\n").writerows(rows)
    return buf.getvalue().encode(encoding)

"""レビューで見つかった問題の再発防止テスト。"""

import io
import logging

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tracetool import colschema, readers
from tracetool.app import create_app
from tracetool.db import Database

from .conftest import make_csv, make_xlsx
from .test_flow import REQ_SCHEMA, do_import, ok, req_csv, scr_schema, scr_xlsx, setup


# --- 取り込み順と要確認 ------------------------------------------------------------


def test_import_lower_first_does_not_make_links_suspect(client):
    req, scr, rel = setup(client)
    do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2)  # 下位を先に取り込む
    do_import(client, req, "req.csv", req_csv())
    ev = ok(client.get(f"/api/trace/{rel}"))
    assert ev["suspect"] == []
    assert len(ev["links"]) == 4


def test_link_to_missing_id_is_not_suspect_when_id_appears(client):
    req, scr, rel = setup(client)
    do_import(client, req, "req.csv", req_csv())
    do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2)  # SCR-04 → REQ-999 はリンク切れ
    ev = ok(client.get(f"/api/trace/{rel}"))
    broken = ev["broken"]
    assert len(broken) == 1
    ok(client.post("/api/links/ack", json={"ids": [broken[0]["id"]]}))  # リンク切れのまま確認済みにしても

    data = make_csv(
        [["要件ID", "要件名"], ["REQ-001", "ログイン"], ["REQ-002", "検索"], ["REQ-003", "出力"], ["REQ-999", "追加"]]
    )
    schema = {"columns": [REQ_SCHEMA["columns"][0], REQ_SCHEMA["columns"][1]]}
    do_import(client, req, "req2.csv", data, schema=schema)
    ev = ok(client.get(f"/api/trace/{rel}"))
    assert ev["broken"] == []
    # REQ-999 が現れたこと自体では要確認にしない（他のリンクは内容が変わったので要確認になる）
    assert all(l["upper_item_id"] != "REQ-999" for l in ev["suspect"])


def test_relation_created_after_imports(client):
    req = ok(client.post("/api/documents", json={"name": "要件", "schema": REQ_SCHEMA}))["id"]
    scr = ok(client.post("/api/documents", json={"name": "画面", "schema": {"columns": scr_schema(req)["columns"][:2]}}))["id"]
    do_import(client, req, "req.csv", req_csv())
    do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2, schema={"columns": scr_schema(req)["columns"][:2]})
    # 参照 ID 列を後から追加して再取り込み → 関係が無いのでエラー
    v = do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2, schema=scr_schema(req))
    assert v["error_count"] == 1
    rel = ok(client.post("/api/relations", json={"upper_doc_id": req, "lower_doc_id": scr}))["id"]
    do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2, schema=scr_schema(req))
    ev = ok(client.get(f"/api/trace/{rel}"))
    assert len(ev["links"]) == 4 and ev["suspect"] == []


# --- フラグの組み合わせ ------------------------------------------------------------


def test_manual_and_auto_on_same_link(client):
    req, scr, rel = setup(client)
    do_import(client, req, "req.csv", req_csv())
    do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2)
    # 自動リンク REQ-002 → SCR-02 に手動でも同じリンクを付ける
    lid = ok(client.post("/api/links", json={"relation_id": rel, "upper_item_id": "REQ-002", "lower_item_id": "SCR-02"}))["id"]
    ev = ok(client.get(f"/api/trace/{rel}"))
    assert [l["origin"] for l in ev["links"] if l["id"] == lid] == ["manual"]
    # 削除すると手動も自動も外れる（自動は無効化として残る）
    ok(client.delete(f"/api/links/{lid}"))
    ev = ok(client.get(f"/api/trace/{rel}"))
    assert all(l["id"] != lid for l in ev["links"])
    ok(client.post(f"/api/links/{lid}/restore"))
    ev = ok(client.get(f"/api/trace/{rel}"))
    assert [l["origin"] for l in ev["links"] if l["id"] == lid] == ["auto"]


def test_both_sides_have_ref_columns(client):
    """上位・下位の両方に参照 ID 列がある場合、片方の再取り込みでもう片方のフラグが消えない。"""
    up = ok(client.post("/api/documents", json={"name": "U"}))["id"]
    lo = ok(client.post("/api/documents", json={"name": "L"}))["id"]
    rel = ok(client.post("/api/relations", json={"upper_doc_id": up, "lower_doc_id": lo}))["id"]
    up_schema = {"columns": [{"key": "u", "name": "ID", "type": "id"}, {"key": "ur", "name": "下位", "type": "string", "ref_document_id": lo}]}
    lo_schema = {"columns": [{"key": "l", "name": "ID", "type": "id"}, {"key": "lr", "name": "上位", "type": "string", "ref_document_id": up}]}
    do_import(client, up, "u.csv", make_csv([["ID", "下位"], ["U1", "L1"], ["U2", "L2"]]), schema=up_schema)
    do_import(client, lo, "l.csv", make_csv([["ID", "上位"], ["L1", "U1"], ["L2", ""]]), schema=lo_schema)
    assert len(ok(client.get(f"/api/trace/{rel}"))["links"]) == 2
    # 下位側の参照を消しても、上位側の参照から U1-L1 は残る
    do_import(client, lo, "l.csv", make_csv([["ID", "上位"], ["L1", ""], ["L2", ""]]), schema=lo_schema)
    pairs = {(l["upper_item_id"], l["lower_item_id"]) for l in ok(client.get(f"/api/trace/{rel}"))["links"]}
    assert pairs == {("U1", "L1"), ("U2", "L2")}


def test_item_summary_with_multiple_relations(client):
    a = ok(client.post("/api/documents", json={"name": "A"}))["id"]
    b = ok(client.post("/api/documents", json={"name": "B"}))["id"]
    c = ok(client.post("/api/documents", json={"name": "C"}))["id"]
    r_ab = ok(client.post("/api/relations", json={"upper_doc_id": a, "lower_doc_id": b}))["id"]
    ok(client.post("/api/relations", json={"upper_doc_id": a, "lower_doc_id": c}))
    for doc, ids in ((a, ["A1", "A2"]), (b, ["B1"]), (c, ["C1"])):
        do_import(client, doc, "x.csv", make_csv([["ID"]] + [[i] for i in ids]))
    ok(client.post("/api/links", json={"relation_id": r_ab, "upper_item_id": "A1", "lower_item_id": "B1"}))
    vid = ok(client.get(f"/api/documents/{a}"))["latest_version_id"]
    res = ok(client.get(f"/api/versions/{vid}/items"))
    trace = {i["item_id"]: i["trace"] for i in res["items"]}
    # A1 は B へのリンクはあるが C へのリンクが無い → 下位なし
    assert trace["A1"]["no_lower"] is True and trace["A1"]["lower_count"] == 1
    assert res["has_lower"] is True and res["has_upper"] is False


def test_delete_all_versions_then_reimport(client):
    req, scr, rel = setup(client)
    do_import(client, req, "req.csv", req_csv())
    v = do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2)["version_id"]
    ok(client.delete(f"/api/versions/{v}"))
    assert ok(client.get(f"/api/trace/{rel}"))["links"] == []
    do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2)
    assert len(ok(client.get(f"/api/trace/{rel}"))["links"]) == 4


# --- 文書の削除と ID ------------------------------------------------------------------


def test_document_ids_are_not_reused_and_refs_cleared(client):
    req, scr, rel = setup(client)
    assert ok(client.get(f"/api/documents/{req}"))["referenced_by"] == ["画面定義書"]
    ok(client.delete(f"/api/documents/{req}"))
    new = ok(client.post("/api/documents", json={"name": "新しい文書"}))["id"]
    assert new != req
    cols = ok(client.get(f"/api/documents/{scr}"))["schema"]["columns"]
    assert all(c["ref_document_id"] is None for c in cols)


def test_update_document_without_schema_keeps_schema(client):
    doc = ok(client.post("/api/documents", json={"name": "A", "schema": REQ_SCHEMA}))["id"]
    ok(client.put(f"/api/documents/{doc}", json={"name": "B"}))
    d = ok(client.get(f"/api/documents/{doc}"))
    assert d["name"] == "B" and len(d["schema"]["columns"]) == 5


def test_migration_from_v1(tmp_path):
    import sqlite3

    path = tmp_path / "old.db"
    conn = sqlite3.connect(path)
    conn.executescript(
        """
        CREATE TABLE meta (key TEXT PRIMARY KEY, value TEXT);
        INSERT INTO meta VALUES ('schema_version', '1');
        CREATE TABLE documents (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, description TEXT NOT NULL DEFAULT '',
          schema_json TEXT NOT NULL, sort_order INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL);
        CREATE TABLE versions (id INTEGER PRIMARY KEY, document_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
          version_no INTEGER NOT NULL, label TEXT NOT NULL DEFAULT '', source_filename TEXT NOT NULL, import_settings TEXT NOT NULL,
          schema_json TEXT NOT NULL, row_count INTEGER NOT NULL, imported_at TEXT NOT NULL, UNIQUE (document_id, version_no));
        CREATE TABLE items (version_id INTEGER NOT NULL REFERENCES versions(id) ON DELETE CASCADE, item_id TEXT NOT NULL,
          row_no INTEGER NOT NULL, data_json TEXT NOT NULL, invalid_json TEXT NOT NULL DEFAULT '[]', content_hash TEXT NOT NULL,
          PRIMARY KEY (version_id, item_id));
        CREATE TABLE relations (id INTEGER PRIMARY KEY, upper_doc_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
          lower_doc_id INTEGER NOT NULL REFERENCES documents(id) ON DELETE CASCADE, UNIQUE (upper_doc_id, lower_doc_id));
        CREATE TABLE links (id INTEGER PRIMARY KEY, relation_id INTEGER NOT NULL REFERENCES relations(id) ON DELETE CASCADE,
          upper_item_id TEXT NOT NULL, lower_item_id TEXT NOT NULL, manual INTEGER NOT NULL DEFAULT 0,
          auto_by_upper INTEGER NOT NULL DEFAULT 0, auto_by_lower INTEGER NOT NULL DEFAULT 0, auto_disabled INTEGER NOT NULL DEFAULT 0,
          upper_hash_ack TEXT, lower_hash_ack TEXT, created_at TEXT NOT NULL, UNIQUE (relation_id, upper_item_id, lower_item_id));
        CREATE INDEX idx_links_upper ON links(relation_id, upper_item_id);
        INSERT INTO documents VALUES (1, 'A', '', '{"columns": []}', 1, 't'), (2, 'B', '', '{"columns": []}', 2, 't');
        INSERT INTO versions VALUES (1, 1, 1, '', 'a.csv', '{}', '{"columns": []}', 1, 't');
        INSERT INTO items VALUES (1, 'X', 1, '{}', '[]', 'h');
        INSERT INTO relations VALUES (1, 1, 2);
        INSERT INTO links VALUES (1, 1, 'X', 'Y', 1, 0, 0, 0, 'h', NULL, 't');
        """
    )
    conn.close()
    db = Database(path)
    with db.read() as c:
        assert c.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0] == "2"
        assert "AUTOINCREMENT" in c.execute("SELECT sql FROM sqlite_master WHERE name='documents'").fetchone()[0]
        assert c.execute("SELECT COUNT(*) FROM links").fetchone()[0] == 1
        assert c.execute("SELECT COUNT(*) FROM items").fetchone()[0] == 1
        assert c.execute("PRAGMA foreign_key_check").fetchall() == []
    with db.tx() as c:
        c.execute("DELETE FROM documents WHERE id = 2")
        c.execute("INSERT INTO documents(name, schema_json, created_at) VALUES ('C', '{}', 't')")
        assert c.execute("SELECT id FROM documents WHERE name='C'").fetchone()[0] == 3
        assert c.execute("SELECT COUNT(*) FROM links").fetchone()[0] == 0  # 外部キーのカスケードが効いている
    db.close()


# --- 出力 ---------------------------------------------------------------------------


def _xlsx_rows(content: bytes, sheet_index: int = 0):
    wb = load_workbook(io.BytesIO(content))
    ws = wb.worksheets[sheet_index]
    return ws, [[c.value for c in row] for row in ws.iter_rows()]


def test_export_with_control_chars_and_formulas(client):
    doc_a = ok(client.post("/api/documents", json={"name": "A"}))["id"]
    doc_b = ok(client.post("/api/documents", json={"name": "B"}))["id"]
    rel = ok(client.post("/api/relations", json={"upper_doc_id": doc_a, "lower_doc_id": doc_b}))["id"]
    do_import(client, doc_a, "a.csv", make_csv([["ID", "名前"], ["A1", "危険\x0bな値"], ["=1+1", "@SUM(1)"]]))
    do_import(client, doc_b, "b.csv", make_csv([["ID", "名前"], ["B1", "x"]]))

    r = client.get(f"/api/export/untraced?relation={rel}&format=xlsx")
    assert r.status_code == 200, r.text
    ws, rows = _xlsx_rows(r.content, 1)
    flat = [v for row in rows for v in row]
    assert "危険な値" in flat  # 制御文字は除去
    cell = next(c for row in ws.iter_rows() for c in row if c.value == "=1+1")
    assert cell.data_type == "s"  # 数式ではなく文字列として、値を変えずに書く

    r = client.get(f"/api/export/untraced?relation={rel}&format=csv")
    text = r.content.decode("utf-8-sig")
    assert "'=1+1" in text and "'@SUM(1)" in text


def test_export_matrix_content(client):
    req, scr, rel = setup(client)
    do_import(client, req, "req.csv", req_csv())
    do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2)
    r = client.get(f"/api/export/matrix?relation={rel}&format=xlsx")
    _, rows = _xlsx_rows(r.content)
    assert rows[0][0] == "要件定義書 ID"
    body = rows[1:]
    assert ["REQ-001", "ログイン", "SCR-01", "ログイン画面", "自動", "正常"] in body
    assert ["REQ-003", "出力", None, None, None, "下位なし"] in body
    assert any(r[0] == "REQ-999" and r[5] == "リンク切れ" for r in body)


def test_diff_detects_type_change():
    from tracetool.services.diff import _same

    assert not _same(1, True)
    assert not _same([1], [True])
    assert _same(["a", 1], ["a", 1])


# --- セキュリティ・ログ ----------------------------------------------------------------


def test_host_header_is_checked(db):
    c = TestClient(create_app(db, port=8765), base_url="http://127.0.0.1:8765")
    assert c.get("/api/documents").status_code == 200
    assert c.get("/api/documents", headers={"host": "localhost:8765"}).status_code == 200
    assert c.get("/api/documents", headers={"host": "evil.example.com:8765"}).status_code == 400


def test_origin_is_checked_for_writes(db):
    c = TestClient(create_app(db, port=8765), base_url="http://127.0.0.1:8765")
    body = {"name": "A"}
    assert c.post("/api/documents", json=body, headers={"origin": "http://evil.example.com"}).status_code == 403
    assert c.post("/api/documents", json=body, headers={"origin": "http://127.0.0.1:9999"}).status_code == 403
    assert c.post("/api/documents", json=body, headers={"sec-fetch-site": "cross-site"}).status_code == 403
    assert c.post("/api/documents", json=body, headers={"origin": "http://127.0.0.1:8765"}).status_code == 200


def test_unexpected_error_log_has_no_values(client, caplog, monkeypatch):
    from tracetool.services import export as export_svc

    def boom(*args, **kwargs):
        raise ValueError("CONFIDENTIAL-VALUE")

    monkeypatch.setattr(export_svc, "untraced_sheets", boom)
    with caplog.at_level(logging.ERROR, logger="tracetool"):
        r = client.get("/api/export/untraced?format=xlsx")
    assert r.status_code == 500
    assert "CONFIDENTIAL-VALUE" not in r.text
    assert "CONFIDENTIAL-VALUE" not in caplog.text
    assert "ValueError" in caplog.text


def test_invalid_body_does_not_echo_values(client):
    r = client.post("/api/links", json={"relation_id": "SECRET-ID", "upper_item_id": "a", "lower_item_id": "b"})
    assert r.status_code == 422
    assert "SECRET-ID" not in r.text and "relation_id" in r.text


# --- 読み込み・正規化 ------------------------------------------------------------------


def test_header_rename_does_not_collide():
    assert readers.build_table([["A", "A", "A(2)"]], 1).headers == ["A", "A(3)", "A(2)"]


def test_large_int_precision():
    col = colschema.normalize_schema({"columns": [{"key": "k", "name": "n", "type": "int"}]})["columns"][0]
    assert colschema.normalize_cell(col, "12345678901234567890") == (12345678901234567890, False)
    assert colschema.normalize_cell(col, "12345678901234567890.00") == (12345678901234567890, False)


def test_tab_and_space_delimiters():
    assert colschema.split_list("A\tB C", ["\t", " "]) == ["A", "B", "C"]


def test_missing_ref_marked_invalid(client):
    req, scr, rel = setup(client)
    do_import(client, req, "req.csv", req_csv())
    vid = do_import(client, scr, "scr.xlsx", scr_xlsx(), header_row=2)["version_id"]
    res = ok(client.get(f"/api/versions/{vid}/items", params={"f.__invalid": "1"}))
    assert [i["item_id"] for i in res["items"]] == ["SCR-04"]


def test_xlsx_formatted_far_rows_are_ignored():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.append(["ID", "名前"])
    ws.append(["A1", "x"])
    ws.cell(row=50000, column=30).font = Font(bold=True)  # 書式だけのセル
    buf = io.BytesIO()
    wb.save(buf)
    grid = readers.read_sheet(readers.open_workbook(buf.getvalue()), ws.title)
    assert grid == [["ID", "名前"], ["A1", "x"]]


def test_tsv_and_manual_encoding(client):
    doc = ok(client.post("/api/documents", json={"name": "T"}))["id"]
    data = make_csv([["ID", "名前"], ["T1", "日本語"]], encoding="cp932", delimiter="\t")
    start = ok(client.post("/api/imports", data={"document_id": doc}, files={"file": ("t.tsv", data)}))
    sid = start["session_id"]
    s = ok(client.put(f"/api/imports/{sid}/settings", json={"header_row": 1, "encoding": "cp932"}))
    assert s["headers"] == ["ID", "名前"] and s["encoding"] == "cp932"
    r = client.put(f"/api/imports/{sid}/settings", json={"header_row": 1, "encoding": "utf-8"})
    assert r.status_code == 400  # UTF-8 としては読めない


def test_date_cells_are_iso():
    import datetime as dt

    data = make_xlsx({"S": [["ID", "日付"], ["A", dt.datetime(2026, 10, 3)], ["B", dt.datetime(2026, 10, 3, 9, 30)]]})
    grid = readers.read_sheet(readers.open_workbook(data), "S")
    assert grid[1][1] == "2026-10-03" and grid[2][1] == "2026-10-03 09:30:00"


# --- テキストの貼り付け・使われている値から enum ------------------------------------


def test_import_from_pasted_tsv_and_csv(client):
    doc = ok(client.post("/api/documents", json={"name": "P"}))["id"]
    # Excel からコピーした表（タブ区切り、セル内改行は引用符付き）
    text = 'ID\t名前\r\nP1\t"1行目\n2行目"\r\nP2\tb\r\n'
    start = ok(client.post("/api/imports/text", json={"document_id": doc, "text": text}))
    assert start["format"] == "tsv"
    s = ok(client.put(f"/api/imports/{start['session_id']}/settings", json={"header_row": 1}))
    assert s["headers"] == ["ID", "名前"] and s["row_count"] == 2
    v = ok(client.put(f"/api/imports/{start['session_id']}/validate", json={"schema": s["schema"]}))
    assert v["preview"][0]["data"][s["schema"]["columns"][1]["key"]] == "1行目\n2行目"
    ok(client.post(f"/api/imports/{start['session_id']}/commit", json={"schema": v["schema"]}))

    start = ok(client.post("/api/imports/text", json={"document_id": doc, "text": "ID,名前\nP3,c\n"}))
    assert start["format"] == "csv"
    r = client.post("/api/imports/text", json={"document_id": doc, "text": "   "})
    assert r.status_code == 400


def test_distinct_values_for_enum(client):
    doc = ok(client.post("/api/documents", json={"name": "E"}))["id"]
    data = make_csv([["ID", "分類", "タグ"], ["E1", "機能", "a;b"], ["E2", "非機能", "b"], ["E3", "機能", ""]])
    start = ok(client.post("/api/imports", data={"document_id": doc}, files={"file": ("e.csv", data)}))
    sid = start["session_id"]
    s = ok(client.put(f"/api/imports/{sid}/settings", json={"header_row": 1}))
    cls, tag = s["schema"]["columns"][1], s["schema"]["columns"][2]
    assert ok(client.put(f"/api/imports/{sid}/distinct", json={"column": cls}))["values"] == ["機能", "非機能"]
    tag_list = {**tag, "list": {"delimiters": [";"]}}
    assert ok(client.put(f"/api/imports/{sid}/distinct", json={"column": tag_list}))["values"] == ["a", "b"]
    v = ok(client.put(f"/api/imports/{sid}/validate", json={"schema": s["schema"]}))
    ok(client.post(f"/api/imports/{sid}/commit", json={"schema": v["schema"]}))

    # 取り込み後の文書から（最新版の値）
    assert ok(client.get(f"/api/documents/{doc}/distinct", params={"key": cls["key"]}))["values"] == ["機能", "非機能"]
    assert client.get(f"/api/documents/{doc}/distinct", params={"key": "nope"}).status_code == 400


def test_distinct_values_too_many(client):
    doc = ok(client.post("/api/documents", json={"name": "M"}))["id"]
    data = make_csv([["ID", "名前"]] + [[f"M{i}", f"値{i}"] for i in range(colschema.MAX_ENUM_VALUES + 1)])
    start = ok(client.post("/api/imports", data={"document_id": doc}, files={"file": ("m.csv", data)}))
    s = ok(client.put(f"/api/imports/{start['session_id']}/settings", json={"header_row": 1}))
    r = client.put(f"/api/imports/{start['session_id']}/distinct", json={"column": s["schema"]["columns"][1]})
    assert r.status_code == 400 and "多すぎます" in r.json()["error"]["message"]

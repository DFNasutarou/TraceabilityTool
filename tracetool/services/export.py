"""Excel / CSV 出力。ファイルはメモリ上で作り、サーバ側には保存しない。"""

from __future__ import annotations

import csv
import io
import sqlite3

from .. import colschema
from ..errors import AppError
from . import diff as diff_svc
from . import documents as doc_svc
from . import relations as rel_svc
from . import trace as trace_svc
from . import versions as ver_svc

STATUS_LABEL = {"ok": "正常", "suspect": "要確認", "broken": "リンク切れ"}
ORIGIN_LABEL = {"manual": "手動", "auto": "自動"}


class Sheet:
    def __init__(self, title: str, header: list[str]):
        self.title = title
        self.header = header
        self.rows: list[list] = []


def item_labels(conn: sqlite3.Connection, doc_id: int) -> dict[str, str]:
    v = ver_svc.latest_version(conn, doc_id)
    if v is None:
        return {}
    disp = colschema.display_column_key(v["schema"])
    return {
        i["item_id"]: colschema.as_text(i["data"].get(disp)) if disp else ""
        for i in ver_svc.load_items(conn, v["id"])
    }


def matrix_sheets(conn: sqlite3.Connection, rel_id: int) -> list[Sheet]:
    ev = trace_svc.evaluate_relation(conn, rel_id)
    rel = ev["relation"]
    up_name = doc_svc.get_document(conn, rel["upper_doc_id"])["name"]
    lo_name = doc_svc.get_document(conn, rel["lower_doc_id"])["name"]
    up_labels = item_labels(conn, rel["upper_doc_id"])
    lo_labels = item_labels(conn, rel["lower_doc_id"])
    sheet = Sheet(
        f"{up_name}→{lo_name}",
        [f"{up_name} ID", f"{up_name} 表示列", f"{lo_name} ID", f"{lo_name} 表示列", "生成元", "状態"],
    )
    for link in sorted(ev["links"], key=lambda l: (l["upper_item_id"], l["lower_item_id"])):
        sheet.rows.append(
            [
                link["upper_item_id"],
                up_labels.get(link["upper_item_id"], ""),
                link["lower_item_id"],
                lo_labels.get(link["lower_item_id"], ""),
                ORIGIN_LABEL[link["origin"]],
                STATUS_LABEL[link["status"]],
            ]
        )
    for item_id in ev["upper"]["untraced"]:
        sheet.rows.append([item_id, up_labels.get(item_id, ""), "", "", "", "下位なし"])
    for item_id in ev["lower"]["untraced"]:
        sheet.rows.append(["", "", item_id, lo_labels.get(item_id, ""), "", "上位なし"])
    return [sheet]


def untraced_sheets(conn: sqlite3.Connection, rel_id: int | None) -> list[Sheet]:
    rel_ids = [rel_id] if rel_id else [r["id"] for r in rel_svc.list_relations(conn)]
    summary = Sheet("集計", ["関係", "上位項目数", "下位なし", "上位網羅率", "下位項目数", "上位なし", "下位網羅率", "リンク切れ", "要確認"])
    detail = Sheet("一覧", ["関係", "区分", "文書", "ID", "表示列"])
    for rid in rel_ids:
        ev = trace_svc.evaluate_relation(conn, rid)
        rel = ev["relation"]
        up_name = doc_svc.get_document(conn, rel["upper_doc_id"])["name"]
        lo_name = doc_svc.get_document(conn, rel["lower_doc_id"])["name"]
        rel_label = f"{up_name}→{lo_name}"
        up_labels = item_labels(conn, rel["upper_doc_id"])
        lo_labels = item_labels(conn, rel["lower_doc_id"])

        def pct(v):
            return "" if v is None else f"{v * 100:.1f}%"

        summary.rows.append(
            [
                rel_label,
                ev["upper"]["total"], len(ev["upper"]["untraced"]), pct(ev["upper"]["coverage"]),
                ev["lower"]["total"], len(ev["lower"]["untraced"]), pct(ev["lower"]["coverage"]),
                len(ev["broken"]), len(ev["suspect"]),
            ]
        )
        for i in ev["upper"]["untraced"]:
            detail.rows.append([rel_label, "下位なし", up_name, i, up_labels.get(i, "")])
        for i in ev["lower"]["untraced"]:
            detail.rows.append([rel_label, "上位なし", lo_name, i, lo_labels.get(i, "")])
        for l in ev["broken"]:
            missing_up = l["upper_item_id"] not in up_labels
            detail.rows.append(
                [rel_label, "リンク切れ", up_name if missing_up else lo_name,
                 l["upper_item_id"] if missing_up else l["lower_item_id"],
                 f"{l['upper_item_id']} → {l['lower_item_id']}"]
            )
        for l in ev["suspect"]:
            detail.rows.append([rel_label, "要確認", "", f"{l['upper_item_id']} → {l['lower_item_id']}", ""])
    return [summary, detail]


def diff_sheets(conn: sqlite3.Connection, from_vid: int, to_vid: int) -> list[Sheet]:
    d = diff_svc.diff_versions(conn, from_vid, to_vid)
    names = d["column_names"]
    sheet = Sheet(f"差分 v{d['from']['version_no']}→v{d['to']['version_no']}", ["区分", "ID", "列名", "変更前", "変更後"])
    kind_label = {"added": "列追加", "removed": "列削除", "renamed": "列名変更", "type_changed": "列の型変更"}
    for c in d["columns"]:
        before = c.get("old_name") or c.get("old_type") or ""
        after = c.get("new_type") or (c["name"] if c["kind"] != "removed" else "")
        sheet.rows.append([kind_label[c["kind"]], "", c["name"], before, after])
    for it in d["added"]:
        sheet.rows.append(["追加", it["item_id"], "", "", ""])
    for it in d["removed"]:
        sheet.rows.append(["削除", it["item_id"], "", "", ""])
    for it in d["changed"]:
        for cell in it["cells"]:
            sheet.rows.append(
                ["変更", it["item_id"], names.get(cell["key"], cell["key"]),
                 colschema.as_text(cell["old"]), colschema.as_text(cell["new"])]
            )
    return [sheet]


def _safe_sheet_title(title: str, used: set[str]) -> str:
    for ch in '[]:*?/\\':
        title = title.replace(ch, "_")
    title = title[:31] or "Sheet"
    base, n = title, 2
    while title in used:
        suffix = f"({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def to_xlsx(sheets: list[Sheet]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for s in sheets:
        ws = wb.create_sheet(_safe_sheet_title(s.title, used))
        ws.append(s.header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDE6F0")
        for row in s.rows:
            # 先頭が = の値を数式として解釈させない
            ws.append([("'" + v) if isinstance(v, str) and v.startswith("=") else v for v in row])
        ws.freeze_panes = "A2"
        for i, _ in enumerate(s.header, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(sheets: list[Sheet]) -> bytes:
    """複数シートは先頭に「シート」列を付けて 1 ファイルにまとめる。"""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    if len(sheets) == 1:
        w.writerow(sheets[0].header)
        w.writerows(sheets[0].rows)
    else:
        for s in sheets:
            w.writerow(["シート"] + s.header)
            w.writerows([[s.title] + r for r in s.rows])
            w.writerow([])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def render(sheets: list[Sheet], fmt: str) -> tuple[bytes, str, str]:
    if fmt == "xlsx":
        return to_xlsx(sheets), "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "csv":
        return to_csv(sheets), "csv", "text/csv; charset=utf-8"
    raise AppError("format は xlsx か csv を指定してください")

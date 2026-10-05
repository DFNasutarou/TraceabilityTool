"""CSV / TSV / XLSX を「行ごとの文字列配列」に変換する。"""

from __future__ import annotations

import csv
import datetime as dt
import io
from dataclasses import dataclass, field
from pathlib import PurePath

from .errors import AppError

ENCODINGS = ("utf-8-sig", "utf-8", "cp932")


@dataclass
class Table:
    """ヘッダと、データ行（元ファイルの行番号付き）。"""

    headers: list[str]
    rows: list[tuple[str, int, list[str]]] = field(default_factory=list)  # (シート名, 行番号, 値)


def detect_format(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext in (".csv", ".txt"):
        return "csv"
    if ext == ".tsv":
        return "tsv"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext == ".xls":
        raise AppError("旧形式の Excel（.xls）には対応していません。.xlsx で保存し直してください")
    raise AppError(f"対応していないファイル形式です: {ext or '(拡張子なし)'}")


def decode_text(data: bytes, encoding: str | None) -> tuple[str, str]:
    """文字コードを指定または自動判定してデコードする。戻り値は (テキスト, 使った文字コード)。"""
    candidates = (encoding,) if encoding else ENCODINGS
    for enc in candidates:
        try:
            text = data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        if enc == "utf-8" and text.startswith("﻿"):
            text = text[1:]
        return text, enc
    if encoding:
        raise AppError(f"文字コード {encoding} として読み込めませんでした")
    raise AppError("文字コードを判定できませんでした（UTF-8 / Shift_JIS 以外の可能性があります）")


def _newlines(text: str) -> str:
    """セル内の改行を LF にそろえる（Windows で作った CSV は CRLF になるため。enum の選択肢との比較で区別しない）。"""
    return text.replace("\r\n", "\n").replace("\r", "\n") if "\r" in text else text


def read_delimited(text: str, delimiter: str) -> list[list[str]]:
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    return [[_newlines(c) for c in r] for r in reader]


def cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return _newlines(str(value))


def open_workbook(data: bytes):
    import openpyxl

    try:
        # 結合セルの情報が必要なので read_only にはしない
        return openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl は壊れたファイルで様々な例外を出す
        raise AppError(f"Excel ファイルを開けませんでした: {type(exc).__name__}") from exc


def read_sheet(wb, sheet_name: str) -> list[list[str]]:
    if sheet_name not in wb.sheetnames:
        raise AppError(f"シート「{sheet_name}」がありません")
    ws = wb[sheet_name]
    # 書式だけが下の方まで設定されたシートでは max_row が非常に大きくなるため、
    # 値の入っているセルと結合範囲から、実際に読む範囲を決める
    max_r = max_c = 0
    for (r, c), cell in ws._cells.items():
        if cell.value is not None and cell.value != "":
            max_r, max_c = max(max_r, r), max(max_c, c)
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= max_r:
            max_c = max(max_c, rng.max_col)
    if max_r == 0:
        return []
    grid = [
        [cell_to_str(v) for v in row]
        for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=max_c, values_only=True)
    ]
    # 結合セルは範囲内の全セルに左上の値を展開する
    for rng in ws.merged_cells.ranges:
        top = grid[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(grid) else ""
        for r in range(rng.min_row - 1, rng.max_row):
            while r >= len(grid):
                grid.append([])
            row = grid[r]
            while len(row) < rng.max_col:
                row.append("")
            for c in range(rng.min_col - 1, rng.max_col):
                row[c] = top
    return grid


def suggest_header_row(grid: list[list[str]], scan_rows: int = 30) -> int:
    """見出しの行を推定する（1 始まり）。

    先頭 scan_rows 行のうち、空でないセルが最も多い行の中で最初の行を見出しとみなす。
    タイトル行や空行が上にある帳票（1 行目タイトル、3 行目見出し など）に対応するため。
    """
    best_row, best_count = 1, 0
    for i, row in enumerate(grid[:scan_rows]):
        count = sum(1 for v in row if str(v).strip())
        if count > best_count:
            best_row, best_count = i + 1, count
    return best_row


def build_table(
    grid: list[list[str]],
    header_row: int,
    sheet_name: str = "",
    data_start: int | None = None,
    data_end: int | None = None,
    stop_at_blank: bool = False,
) -> Table:
    """行番号はすべて 1 始まり。

    data_start: データの最初の行（省略時はヘッダ行の次）。見出しとデータの間に説明の行がある帳票用。
    data_end: データの最後の行（省略時は最後まで）。表の下に注記などがある帳票用。
    stop_at_blank: データの途中で空行が現れたら、そこで終わりとみなす（シートごとに表の長さが違う場合用）。
    """
    if header_row < 1:
        raise AppError("ヘッダ行は 1 以上を指定してください")
    start = data_start if data_start else header_row + 1
    if start <= header_row:
        raise AppError(f"データの開始行（{start} 行目）はヘッダ行（{header_row} 行目）より後にしてください")
    if data_end and data_end < start:
        raise AppError(f"データの終了行（{data_end} 行目）が開始行（{start} 行目）より前です")
    if header_row > len(grid):
        raise AppError(f"ヘッダ行 {header_row} 行目がありません（全 {len(grid)} 行）")
    headers = [h.strip() for h in grid[header_row - 1]]
    while headers and headers[-1] == "":
        headers.pop()
    if not headers:
        raise AppError(f"ヘッダ行 {header_row} 行目が空です")
    # 空のヘッダには名前を付け、重複する名前には (2), (3)… を付ける。
    # 付け替えた名前が元から存在する名前と衝突しないよう、すべての名前を予約してから付ける
    names = [h or f"列{i + 1}" for i, h in enumerate(headers)]
    used: set[str] = set()
    for i, name in enumerate(names):
        if name in used:
            n = 2
            while f"{name}({n})" in used or f"{name}({n})" in names[i + 1 :]:
                n += 1
            name = f"{name}({n})"
        used.add(name)
        headers[i] = name
    table = Table(headers=headers)
    width = len(headers)
    end = min(len(grid), data_end) if data_end else len(grid)
    started = False
    for idx in range(start - 1, end):
        values = (grid[idx] + [""] * width)[:width]
        if all(v.strip() == "" for v in values):
            # 先頭側の空行は読み飛ばし、データが始まった後の空行で終える
            if stop_at_blank and started:
                break
            continue
        started = True
        table.rows.append((sheet_name, idx + 1, values))
    return table


def merge_tables(tables: list[tuple[str, Table]]) -> Table:
    """複数シートの表を連結する。ヘッダが一致しなければエラー。"""
    first_name, first = tables[0]
    merged = Table(headers=list(first.headers))
    for name, t in tables:
        if t.headers != first.headers:
            raise AppError(f"シート「{name}」のヘッダがシート「{first_name}」と一致しません")
        merged.rows.extend(t.rows)
    return merged
